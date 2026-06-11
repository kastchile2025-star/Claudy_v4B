"""Claudy features.documents — Creación y análisis de documentos (refactor v5).

Orquesta la generación de .docx/.pdf/.xlsx/.pptx con contenido del LLM
(la escritura física la hace claudy_powers) y el análisis de documentos
locales (lectura de texto de docx/pptx/pdf, análisis con el LLM).

Se usa como mixin: ClawdPet hereda de DocumentsMixin. Depende de helpers
de la clase compuesta: _llm_structured, _extract_json, _web_context_for_topic,
_find_local_files, _extract_attachment_text, _save_memory, send_quick_message.
"""
import datetime
import os
import re
import subprocess
import sys


class DocumentsMixin:
    def _read_docx_text(self, path):
        try:
            from docx import Document
        except ImportError:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "python-docx"])
            from docx import Document
        doc = Document(path)
        parts = []
        for p in doc.paragraphs:
            txt = (p.text or "").strip()
            if txt:
                parts.append(txt)
        for idx, table in enumerate(doc.tables[:20], 1):
            rows = []
            for row in table.rows[:60]:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                parts.append(f"\nTabla {idx}\n" + "\n".join(rows))
        return "\n".join(parts)

    def _read_pptx_text(self, path):
        try:
            from pptx import Presentation
        except ImportError:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "python-pptx"])
            from pptx import Presentation
        prs = Presentation(path)
        parts = []
        for slide_idx, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                try:
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                            if any(cells):
                                texts.append(" | ".join(cells))
                    elif hasattr(shape, "text"):
                        txt = (shape.text or "").strip()
                        if txt:
                            texts.append(txt)
                except Exception:
                    pass
            if texts:
                parts.append(f"Diapositiva {slide_idx}\n" + "\n".join(texts))
        return "\n\n".join(parts)

    def _read_pdf_text(self, path):
        try:
            from pypdf import PdfReader
        except ImportError:
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "pypdf"])
                from pypdf import PdfReader
            except Exception:
                try:
                    from PyPDF2 import PdfReader
                except Exception:
                    return "[Para leer PDFs instala: pip install pypdf]"
        except Exception:
            try:
                from PyPDF2 import PdfReader
            except Exception:
                return "[Para leer PDFs instala: pip install pypdf]"
        reader = PdfReader(path)
        parts = []
        for page in reader.pages[:40]:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                pass
        return "\n".join(parts).strip()

    def _try_handle_create_document(self, prompt):
        """Detect 'crea(me) (un) archivo X.docx/X.pdf [en CARPETA] con/sobre/del TEMA'
        and create the document with LLM-generated content. Returns (handled, result)."""
        text = prompt.strip()
        # Variant A: "crea(me) archivo|presentacion|excel docx|pdf|xlsx|pptx [NOMBRE] [en LOC] (con|sobre|...) TEMA"
        # Variant B: "crea(me) archivo NOMBRE.docx|.pdf|.xlsx|.pptx [en LOC] (con|sobre|...) TEMA"
        verb_re = r"(?:cr[eé]a(?:me|r)?|genera(?:me)?|haz(?:me)?|hacer|arma(?:me)?|prepara(?:me)?)"
        kind_re = r"(?:archivo|documento|doc|file|presentaci[oó]n|presentacion|diapositivas|slides|deck|excel|hoja(?:\s+de\s+c[aá]lculo)?|planilla|spreadsheet|libro)"
        connector_re = r"(?:con|sobre|del|de|acerca\s+de|que\s+(?:tenga|contenga|diga|incluya|muestre))"
        location_re = r"(?:en|dentro\s+de)"
        fmt_alts = r"(docx|pdf|word|xlsx|excel|pptx|powerpoint|ppt)"

        # Detect format hint also from kind word itself (excel/presentacion)
        kind_match = re.match(rf"^{verb_re}\s+(?:un[ao]?\s+|el\s+|la\s+)?({kind_re})\b", text, re.I)
        kind_word = kind_match.group(1).lower() if kind_match else ""

        # Variant A
        m = re.match(
            rf"^{verb_re}\s+(?:un[ao]?\s+|el\s+|la\s+)?{kind_re}\s+"
            rf"{fmt_alts}\b\s*"
            rf"(?:(?!{location_re}\s|{connector_re}\s)(\S.*?)\s+)?"
            rf"(?:{location_re}\s+(.+?)\s+)?"
            rf"{connector_re}\s+(.+)$",
            text, re.I,
        )
        if m:
            fmt_word = m.group(1)
            name_part = m.group(2) or ""
            location_part = m.group(3)
            topic = m.group(4)
        else:
            # Variant B: name has explicit extension
            m = re.match(
                rf"^{verb_re}\s+(?:un[ao]?\s+|el\s+|la\s+)?{kind_re}\s+"
                rf"(\S+\.(?:docx|pdf|xlsx|pptx))"
                rf"(?:\s+{location_re}\s+(.+?))?"
                rf"\s+{connector_re}\s+(.+)$",
                text, re.I,
            )
            if m:
                fmt_word = None
                name_part = m.group(1)
                location_part = m.group(2)
                topic = m.group(3)
            else:
                # Variant C: format implied by kind word (excel/presentacion sin docx/pdf)
                m = re.match(
                    rf"^{verb_re}\s+(?:un[ao]?\s+|el\s+|la\s+)?{kind_re}\s+"
                    rf"(?:(?!{location_re}\s|{connector_re}\s)(\S.*?)\s+)?"
                    rf"(?:{location_re}\s+(.+?)\s+)?"
                    rf"{connector_re}\s+(.+)$",
                    text, re.I,
                )
                if not m or not kind_word:
                    return False, ""
                fmt_word = None
                name_part = m.group(1) or ""
                location_part = m.group(2)
                topic = m.group(3)

        name_clean = (name_part or "").strip().strip('"\'')
        fmt = ""
        if fmt_word:
            fw = fmt_word.lower()
            if fw in ("docx", "word"):
                fmt = "docx"
            elif fw == "pdf":
                fmt = "pdf"
            elif fw in ("xlsx", "excel"):
                fmt = "xlsx"
            elif fw in ("pptx", "powerpoint", "ppt"):
                fmt = "pptx"
        elif name_clean.lower().endswith(".docx"):
            fmt = "docx"
        elif name_clean.lower().endswith(".pdf"):
            fmt = "pdf"
        elif name_clean.lower().endswith(".xlsx"):
            fmt = "xlsx"
        elif name_clean.lower().endswith(".pptx"):
            fmt = "pptx"
        elif kind_word:
            if "excel" in kind_word or "hoja" in kind_word or "planilla" in kind_word or "spreadsheet" in kind_word or "libro" in kind_word:
                fmt = "xlsx"
            elif "presentaci" in kind_word or "diapositiv" in kind_word or "slides" in kind_word or "deck" in kind_word:
                fmt = "pptx"
        if not fmt:
            return False, ""

        # Resolve location: explicit phrase, or 'esta carpeta' reference, or Downloads default
        folder = ""
        if location_part:
            ref = self._resolve_last_folder_reference(location_part)
            if ref:
                folder = ref
            else:
                low = location_part.lower()
                if any(w in low for w in ("esta carpeta", "ultima carpeta", "última carpeta", "carpeta anterior", "esa carpeta", "misma carpeta")):
                    # Reference but no memory available -> fall back to Downloads silently
                    folder = ""
                else:
                    folder = location_part.strip()
        else:
            ref = self._resolve_last_folder_reference(prompt)
            if ref:
                folder = ref
        if not folder:
            folder = os.path.expanduser("~/Downloads")

        # Strip extension from name to rebuild it cleanly; derive from topic if missing
        bare = re.sub(r"\.(docx|pdf)$", "", name_clean, flags=re.I).strip()
        if not bare:
            slug = re.sub(r"[^A-Za-z0-9_\- ]+", "", topic).strip()
            slug = re.sub(r"\s+", "_", slug)[:60] or "documento"
            bare = slug
        full_path = os.path.join(folder, f"{bare}.{fmt}")

        topic_clean = topic.strip().rstrip(".")

        try:
            import claudy_powers as cp
        except Exception as e:
            return True, f"Error: {e}"

        # ALL formats: try to ground content with web search first
        try:
            self.after(0, lambda: self._set_response_text(f"Buscando datos en internet sobre: {topic_clean}..."))
        except Exception:
            pass
        web_ctx = self._web_context_for_topic(topic_clean, max_results=8)

        # Load professional formatting rules from skill file
        _skill_rules = ""
        try:
            _skill_path = os.path.join(
                os.path.dirname(os.path.abspath(__file__)),
                "..", "..", "skills", "professional-document-writer", "SKILL.md"
            )
            if os.path.isfile(_skill_path):
                with open(_skill_path, "r", encoding="utf-8") as _sf:
                    _skill_rules = _sf.read()
        except Exception:
            pass
        if not _skill_rules:
            _skill_rules = (
                "FORMATO PROFESIONAL:\n"
                "- Título con # TÍTULO. Introducción obligatoria.\n"
                "- Secciones con ## y subsecciones con ###. Nunca saltes niveles.\n"
                "- Usa **negritas** para conceptos clave, *cursivas* para énfasis secundario.\n"
                "- Datos comparativos en TABLAS Markdown (| Col | Col |).\n"
                "- Citas importantes con > formato.\n"
                "- Termina con ## Conclusiones y ## Referencias.\n"
                "- NO uses placeholders como [Insertar aquí]. Escribe contenido real.\n"
                "- NO incluyas saludos ni comentarios fuera del documento."
            )

        # Check if user wants images in natural language document creation
        _wants_images_nl = False
        _img_keywords = ("imagen", "imágenes", "imagenes", "foto", "fotos",
                         "ilustra", "ilustracion", "ilustraciones", "visual",
                         "gráfico", "grafico", "con imagen", "con fotos")
        if any(k in prompt.lower() for k in _img_keywords):
            _wants_images_nl = True

        # Format-specific LLM prompts
        if fmt in ("docx", "pdf"):
            llm_prompt = (
                (web_ctx + "\n\n" if web_ctx else "") +
                f"Escribe un documento profesional en espanol sobre: {topic_clean}.\n"
                + ("USA la informacion de internet de arriba como base factual. Cita fuentes al final cuando uses datos especificos. " if web_ctx else "") +
                f"\nDIRECTRICES DE FORMATO PROFESIONAL (sigue TODAS estas reglas):\n{_skill_rules}\n\n"
                "REGLAS ADICIONALES:\n"
                "- Escribe contenido COMPLETO y REAL. NO uses placeholders como [Introducción], [Sección], [Insertar aquí], etc.\n"
                "- Cada sección debe tener párrafos sustanciales (mínimo 2-3 párrafos).\n"
                "- Incluye al menos UNA tabla Markdown con datos relevantes.\n"
                "- Extensión: 500-900 palabras. Solo el contenido del documento, sin preambulo ni cierre."
            )
            content = self._llm_structured(llm_prompt, timeout=180)
            if not content:
                return True, f"No pude generar contenido para '{topic_clean}'."

            # Download and attach images if requested
            if _wants_images_nl and fmt == "docx":
                try:
                    self.after(0, lambda: self._set_response_text(f"Descargando imágenes para: {topic_clean}..."))
                    image_paths = self._download_report_images(topic_clean, max_images=3)
                except Exception:
                    image_paths = []
                if image_paths:
                    return True, cp.create_docx_with_images(full_path, content, image_paths)

            if fmt == "docx":
                return True, cp.create_docx(full_path, content)
            return True, cp.create_pdf(full_path, content)

        if fmt == "xlsx":
            llm_prompt = (
                (web_ctx + "\n\n" if web_ctx else "") +
                f"Genera datos de hoja de calculo profesional en espanol sobre: {topic_clean}.\n"
                + ("USA los datos de internet de arriba como base. Si los datos son incompletos, completa con cifras plausibles pero MARCA esas filas con '(estimado)' al final. " if web_ctx else "") +
                "Responde SOLO con un JSON valido (sin ```fences```) con esta forma exacta:\n"
                '{"name": "NombreHoja", "headers": ["Col1","Col2",...], "rows": [["v1","v2"],...]}\n'
                "Si el tema requiere varias hojas, usa: {\"sheets\": [{\"name\":..., \"headers\":..., \"rows\":...}, ...]}.\n"
                "Entre 8 y 20 filas con datos realistas y concretos. Incluye una columna 'Fuente' al final si usaste datos web. "
                "Sin texto fuera del JSON."
            )
            raw = self._llm_structured(llm_prompt, timeout=180, expect_json=True)
            data = self._extract_json(raw)
            if not data:
                # Last-ditch: build a sheet from the web search snippets themselves
                lines = [ln.strip() for ln in (raw or "").splitlines() if ln.strip() and not ln.strip().startswith("{")]
                rows = []
                if web_ctx:
                    for line in web_ctx.split("\n"):
                        s = line.strip().lstrip("- ").strip()
                        if s and not s.startswith("DATOS DE INTERNET") and not s.startswith("Fuente:"):
                            rows.append([s[:300]])
                if not rows and lines:
                    rows = [[ln[:300]] for ln in lines[:30]]
                if not rows:
                    return True, f"No pude generar datos para '{topic_clean}'. Intenta de nuevo o se mas especifico."
                data = {"name": (bare[:31] or "Datos"), "headers": ["Informacion"], "rows": rows[:40]}
            return True, cp.create_xlsx(full_path, data, title=bare[:31] or "Datos")

        if fmt == "pptx":
            llm_prompt = (
                (web_ctx + "\n\n" if web_ctx else "") +
                f"Genera una presentacion profesional en espanol sobre: {topic_clean}.\n"
                + ("USA los datos de internet de arriba como base factual. " if web_ctx else "") +
                "Responde SOLO con un JSON valido con esta forma exacta:\n"
                "{\n"
                '  "title": "Titulo principal",\n'
                '  "subtitle": "Subtitulo descriptivo",\n'
                '  "theme": "history|nature|tech|business|education|warm|dark|minimal",\n'
                '  "slides": [\n'
                '    {"title": "Titulo slide", "bullets": ["punto 1","punto 2","punto 3"], "image_query": "descripcion visual en INGLES para generar imagen"},\n'
                "    ...\n"
                "  ]\n"
                "}\n"
                "Reglas:\n"
                "- 6 a 10 slides. Cada slide con 3-5 bullets concisos (max 14 palabras).\n"
                "- `theme`: elige el que mejor encaje con el tema (history para historia, tech para tecnologia, nature para naturaleza, etc.).\n"
                "- `image_query`: SIEMPRE en INGLES, visual concreto y especifico, sin texto en la imagen. Ej: 'mapuche warriors traditional dress 19th century', 'santiago chile colonial architecture'.\n"
                "- Sin texto fuera del JSON. Sin ```fences```."
            )
            raw = self._llm_structured(llm_prompt, timeout=180)
            data = self._extract_json(raw) or {}
            if not data.get("slides"):
                return True, f"No pude generar la estructura para '{topic_clean}'. Intenta otra vez o se mas especifico."
            return True, cp.create_pptx(
                full_path,
                slides=data.get("slides") or [],
                title=data.get("title") or topic_clean,
                subtitle=data.get("subtitle") or "",
                theme=(data.get("theme") or "business").lower(),
                images=True,
            )

        return False, ""

    def _try_local_file_action(self, prompt, lower):
        """Si el usuario indica una UBICACIÓN local + buscar/analizar un archivo:
        ubica el archivo por semejanza, lo analiza y devuelve un mensaje con el resumen.
        Devuelve (handled, result)."""
        if not any(kw in lower for kw in self._LOCAL_LOC_KWS):
            return False, None
        # 1) Nombre/consulta del archivo
        fp = re.search(r'[\w.\-]{2,}\.\w{1,5}', prompt)
        query = fp.group(0) if fp else ""
        if not query:
            m = re.search(r'(?:archivo|fichero)\s+(?:llamado\s+|que\s+se\s+llama\s+)?(.+)$', lower)
            if m:
                query = m.group(1).strip()
        if not query:
            cleaned = lower
            for w in self._LOCAL_LOC_KWS:
                cleaned = cleaned.replace(w, " ")
            cleaned = re.sub(
                r"\b(busca|buscar|búscalo|buscalo|encuentra|encuéntrame|encuentrame|esta|este|esto|ese|esa|eso|el|la|los|las|en|mi|archivo|fichero)\b",
                " ", cleaned)
            query = re.sub(r"\s+", " ", cleaned).strip()
        # Quitar comandos que se cuelan al final ("... y analízalo", "y dime de qué trata")
        query = re.sub(
            r'\s*\b(y|e|,)?\s*(anal[ií]za\w*|anal[ií]zam\w*|dime|mu[eé]stra\w*|res[uú]m\w*|abre\w*|lee\w*|de\s+qu[eé]\s+(se\s+)?trata|qu[eé]\s+trata|por\s+favor)\b.*$',
            '', query, flags=re.IGNORECASE)
        query = query.strip(" .,:;¿?¡!y").strip()
        if (not query or query in ("esta", "este", "esto", "ese", "esa", "eso")):
            last = (getattr(self, "_last_file_query", "") or "").strip()
            query = last.split()[0] if last else ""
        if not query or len(query) < 2:
            return True, "¿Qué archivo busco? Dime el nombre (o parte de él)."
        self._last_file_query = query

        # 2) Carpeta objetivo
        home = os.path.expanduser("~")
        if "descarga" in lower:
            dirs = [os.path.join(home, "Downloads")]
        elif "escritorio" in lower or "desktop" in lower:
            dirs = [os.path.join(home, "Desktop")]
        elif "documento" in lower:
            dirs = [os.path.join(home, "Documents")]
        else:
            dirs = [os.path.join(home, d) for d in ("Downloads", "Desktop", "Documents")] + [home]

        # 3) Buscar (fuzzy) y analizar el más parecido
        scored = self._find_local_files(query, dirs)
        if not scored:
            return True, (f"No encontré ningún archivo parecido a '{query}' en esa carpeta. "
                          "Prueba con otra parte del nombre.")
        ambiguous = len(scored) > 1 and (scored[0][0] - scored[1][0]) < 0.4
        if ambiguous:
            listado = "\n".join(f"  • {os.path.basename(p)}" for _, p in scored[:6])
            return True, (f"Encontré varios archivos parecidos a '{query}':\n{listado}\n\n"
                          "¿Cuál analizo? (dime el nombre)")
        return True, self._analyze_local_file_sync(scored[0][1])

    def _create_pdf_simple(self, title, content):
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        except ImportError:
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", "reportlab", "--quiet"])
                from reportlab.lib.pagesizes import letter
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            except Exception:
                return "Necesito 'reportlab' para crear PDFs: pip install reportlab"
        out_dir = os.path.join(os.path.expanduser("~"), ".claudy", "pdfs")
        os.makedirs(out_dir, exist_ok=True)
        fname = f"claudy_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        fpath = os.path.join(out_dir, fname)
        try:
            doc = SimpleDocTemplate(fpath, pagesize=letter)
            styles = getSampleStyleSheet()
            story = [Paragraph(title, styles['Title']), Spacer(1, 12)]
            for line in content.split("\n"):
                if line.strip():
                    story.append(Paragraph(line.strip(), styles['Normal']))
                else:
                    story.append(Spacer(1, 6))
            doc.build(story)
            return f"PDF creado: {fpath}"
        except Exception as e:
            return f"Error creando PDF: {e}"

    def _analyze_pdf_text(self, path):
        if not os.path.exists(path):
            return f"No encuentro: {path}"
        try:
            from pypdf import PdfReader
        except ImportError:
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", "pypdf", "--quiet"])
                from pypdf import PdfReader
            except Exception:
                return "Necesito 'pypdf' para leer PDFs: pip install pypdf"
        try:
            reader = PdfReader(path)
            n_pages = len(reader.pages)
            text = ""
            for page in reader.pages[:5]:
                pt = page.extract_text()
                if pt:
                    text += pt + "\n"
            summary = text[:2000] if text else "No se pudo extraer texto."
            return f"PDF: {os.path.basename(path)}\nPáginas: {n_pages}\n\nContenido:\n{summary}"
        except Exception as e:
            return f"Error leyendo PDF: {e}"

    def _analyze_local_file_sync(self, path):
        """Lee, analiza y resume un archivo local; guarda el análisis en memoria
        de agentes + Obsidian. Devuelve el resumen (texto)."""
        name = os.path.basename(path)
        folder = os.path.dirname(path)
        try:
            size_kb = round(os.path.getsize(path) / 1024, 1)
        except Exception:
            size_kb = 0
        ext = os.path.splitext(path)[1].lower()
        image_exts = (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif")
        try:
            if ext in image_exts:
                answer = self._analyze_image_native(
                    path, "Describe esta imagen en español y extrae el texto visible.")
            else:
                text = self._extract_attachment_text(path)
                if not text.strip():
                    return f"📄 Encontré **{name}** en {folder} ({size_kb} KB), pero no pude extraer contenido legible."
                prompt2 = (
                    f"Analiza este archivo y dime de qué trata.\nArchivo: {name}\n\n"
                    "Contenido (puede venir truncado):\n```\n" + text + "\n```\n\n"
                    "Responde en español: (1) de qué trata en 1-2 frases, (2) puntos/datos clave, "
                    "(3) si aplica, montos, fechas o totales relevantes."
                )
                answer = self.send_quick_message(prompt2, _skip_skill_action=True, timeout=120)
        except Exception as e:
            return f"Encontré **{name}** en {folder} pero falló el análisis: {e}"
        try:
            self._save_memory("Usuario", f"Analizar archivo local: {name} ({folder})")
            self._save_memory("Claudy", f"[Análisis de archivo: {name}] {answer}")
        except Exception:
            pass
        return f"📄 **{name}**  ·  {folder}  ·  {size_kb} KB\n\n{answer}"

